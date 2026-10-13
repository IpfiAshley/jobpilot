import re
from pathlib import Path

import fitz
from docx import Document


class CVParseError(ValueError):
    """Raised when a CV cannot be read or has no extractable text."""


KNOWN_SKILLS = {
    "python", "java", "javascript", "typescript", "c#", "c++", "go", "rust", "sql",
    "postgresql", "mysql", "mongodb", "redis", "aws", "azure", "gcp", "docker", "kubernetes",
    "terraform", "linux", "git", "fastapi", "django", "flask", "react", "node.js", " .net",
    "machine learning", "deep learning", "scikit-learn", "pandas", "numpy", "spark", "pytorch",
    "tensorflow", "tableau", "power bi", "looker", "excel", "etl", "airflow", "dbt",
    "cybersecurity", "siem", "network security", "ci/cd", "jenkins", "communication", "leadership",
}
SECTION_ALIASES = {
    "experience": ("experience", "employment", "work history", "career history"),
    "education": ("education", "academic background", "qualifications"),
    "projects": ("projects", "selected projects", "personal projects"),
    "skills": ("skills", "technical skills", "technologies", "competencies"),
    "summary": ("summary", "profile", "professional summary", "objective"),
}


def _read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with fitz.open(path) as pdf:
            text = "\n".join(page.get_text("text") for page in pdf)
    elif suffix == ".docx":
        document = Document(path)
        chunks = [p.text for p in document.paragraphs]
        for table in document.tables:
            chunks.extend(cell.text for row in table.rows for cell in row.cells)
        text = "\n".join(chunks)
    else:
        raise CVParseError("Only PDF and DOCX CV files are supported")
    text = text.replace("\x00", " ").strip()
    if not text:
        raise CVParseError("No text found; scanned PDFs require OCR before parsing")
    if len(text) > 1_000_000:
        raise CVParseError("CV text exceeds the 1,000,000 character processing limit")
    return text


def parse_cv(path: str | Path) -> dict:
    """Extract a reviewable, heuristic profile from a PDF or DOCX file."""
    file_path = Path(path)
    if not file_path.is_file():
        raise CVParseError("CV file does not exist")
    try:
        text = _read_text(file_path)
    except CVParseError:
        raise
    except Exception as exc:
        raise CVParseError("The uploaded file could not be read as a valid PDF or DOCX") from exc
    sections = _split_sections(text)
    skills = _extract_skills(text)
    experience_lines = _extract_bullets(sections.get("experience", ""))
    project_lines = _extract_bullets(sections.get("projects", ""))
    education_lines = _extract_bullets(sections.get("education", ""))
    return {
        "name": _extract_name(text),
        "email": _first_match(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text),
        "phone": _first_match(r"(?:\+?\d[\d ()-]{7,}\d)", text),
        "summary": _clean_section(sections.get("summary", ""))[:2000],
        "skills": skills,
        "experience": experience_lines,
        "education": education_lines,
        "projects": project_lines,
        "parser_notes": ["Heuristic extraction; check section boundaries and extracted skills before use."],
    }


def _split_sections(text: str) -> dict[str, str]:
    lines = text.splitlines()
    labels = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}
    found: list[tuple[int, str]] = []
    for i, line in enumerate(lines):
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if normalized in labels and len(line.strip()) < 60:
            found.append((i, labels[normalized]))
    result: dict[str, str] = {}
    for idx, (line_no, section) in enumerate(found):
        end = found[idx + 1][0] if idx + 1 < len(found) else len(lines)
        result[section] = "\n".join(lines[line_no + 1 : end]).strip()
    return result


def _extract_skills(text: str) -> list[str]:
    lower = text.lower()
    return sorted(skill.strip() for skill in KNOWN_SKILLS if re.search(rf"(?<!\w){re.escape(skill.strip())}(?!\w)", lower))


def _extract_bullets(section: str) -> list[str]:
    items = []
    for line in section.splitlines():
        cleaned = re.sub(r"^\s*[•●▪◦*\-–—]+\s*", "", line).strip()
        if len(cleaned) >= 20:
            items.append(cleaned[:1000])
    return items[:100]


def _clean_section(section: str) -> str:
    return re.sub(r"\s+", " ", section).strip()


def _first_match(pattern: str, text: str) -> str | None:
    match = re.search(pattern, text, flags=re.IGNORECASE)
    return match.group(0).strip() if match else None


def _extract_name(text: str) -> str | None:
    for line in text.splitlines()[:8]:
        candidate = re.sub(r"\s+", " ", line).strip()
        if 2 <= len(candidate) <= 70 and not any(ch.isdigit() for ch in candidate) and "@" not in candidate:
            if len(candidate.split()) <= 5:
                return candidate
    return None
