from io import BytesIO
import re

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape


def _safe_lines(profile: dict, tailored: dict) -> list[tuple[str, list[str]]]:
    name = str(profile.get("name") or "Candidate")[:100]
    contact = " · ".join(str(x) for x in (profile.get("email"), profile.get("phone")) if x)
    sections = [(name, [contact] if contact else [])]
    if tailored.get("summary"):
        sections.append(("Professional Summary", [str(tailored["summary"])]))
    if tailored.get("skills"):
        sections.append(("Skills", [", ".join(map(str, tailored["skills"]))]))
    source_bullets = [*profile.get("experience", []), *profile.get("projects", [])]
    edited_by_index = {item.get("source_index"): str(item.get("text", "")) for item in tailored.get("bullets", []) if isinstance(item, dict)}
    experience = [edited_by_index.get(i, str(value)) for i, value in enumerate(source_bullets)]
    if experience:
        sections.append(("Experience & Projects", experience))
    if profile.get("education"):
        sections.append(("Education", [str(x) for x in profile["education"]]))
    return [(heading[:120], [str(line)[:3000] for line in lines if str(line).strip()]) for heading, lines in sections]


def export_docx(profile: dict, tailored: dict) -> bytes:
    document = Document()
    document.add_heading(str(profile.get("name") or "Candidate")[:100], 0)
    contact = " | ".join(str(x) for x in (profile.get("email"), profile.get("phone")) if x)
    if contact:
        document.add_paragraph(contact)
    for heading, lines in _safe_lines(profile, tailored)[1:]:
        document.add_heading(heading, level=1)
        for line in lines:
            if heading in ("Experience & Projects", "Education"):
                document.add_paragraph(line, style="List Bullet")
            else:
                document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(profile: dict, tailored: dict) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    story = []
    for heading, lines in _safe_lines(profile, tailored):
        story.append(Paragraph(escape(heading), styles["Title"] if heading == str(profile.get("name") or "Candidate")[:100] else styles["Heading2"]))
        for line in lines:
            safe = escape(re.sub(r"\s+", " ", line))
            story.append(Paragraph(safe, styles["BodyText"]))
            story.append(Spacer(1, 4))
        story.append(Spacer(1, 6))
    doc.build(story)
    return output.getvalue()
